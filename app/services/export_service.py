"""Export business logic — Excel export of vote / payment data per event."""
import io
from typing import Optional

from sqlalchemy.ext.asyncio import AsyncSession

from app.repositories import event as event_repo
from app.repositories import payment as payment_repo
from app.repositories import team as team_repo


async def export_vote_results(session: AsyncSession, event_id: str) -> bytes:
    """Return an .xlsx (in-memory bytes) of teams + total votes for an event."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    event = await event_repo.get_event(session, event_id)
    if not event or event.deleted_at is not None:
        raise ValueError("Event tidak ditemukan")

    pairs = await team_repo.get_teams_with_votes(session, event_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hasil Voting"

    ws.append(["Rangking", "Nama Tim", "Asal Sekolah", "Total Vote"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for rank, (team, votes) in enumerate(pairs, start=1):
        ws.append([rank, team.name, team.school or "", votes])

    for col in ("A", "B", "C", "D"):
        ws.column_dimensions[col].width = 24

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def export_payment_report(
    session: AsyncSession,
    event_id: str,
    status: Optional[str] = None,
) -> bytes:
    """Return an .xlsx of payments for an event (filter by status optional)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    from app.enums.payment_status import PaymentStatus

    status_enum = PaymentStatus(status) if status else None
    payments = await payment_repo.list_payments(session, event_id=event_id, status=status_enum, limit=5000)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transaksi"

    ws.append(["Invoice", "Tim", "Paket", "Qty", "Total", "Vote", "Status", "Waktu"])
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for payment in payments:
        team = await team_repo.get_team(session, payment.team_id)
        ws.append(
            [
                payment.id,
                team.name if team else payment.team_id,
                f"{payment.package_label} ({payment.package_code})",
                payment.qty,
                payment.amount,
                payment.votes,
                payment.status.value,
                payment.created_at.replace(tzinfo=None).isoformat(sep=" "),
            ]
        )

    for col in ("A", "B", "C", "D", "E", "F", "G", "H"):
        ws.column_dimensions[col].width = 24

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
